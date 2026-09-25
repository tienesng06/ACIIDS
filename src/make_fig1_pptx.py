"""Create the editable PowerPoint counterpart of manuscript Figure 1.

The slide mirrors sections/fig-framework.tex.  It uses the same genuine SoccerNet
frame and parses the real leave-EPL-out score traces from
generated/fig_framework_data.tex.  All diagram elements and score traces are
native PowerPoint objects; only the source broadcast frame is raster.

Usage: python src/make_fig1_pptx.py [output.pptx]
"""
import os
import re
import sys

from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
from common import ROOT

MANUSCRIPT = os.path.join(os.path.dirname(ROOT), "manuscript")
DATA_TEX = os.path.join(MANUSCRIPT, "generated", "fig_framework_data.tex")
FRAME = os.path.join(
    MANUSCRIPT, "images", "dataset-source",
    "soccernet-liverpool-chelsea-goal-2359.png",
)

INK = RGBColor(35, 55, 75)
VISUAL = RGBColor(0, 114, 178)
AUDIO = RGBColor(213, 94, 0)
GOLD = RGBColor(230, 159, 0)
TEAL = RGBColor(0, 158, 115)
WHITE = RGBColor(255, 255, 255)


def macro(source, name):
    hit = re.search(r"\\newcommand\{\\" + re.escape(name) + r"\}\{(.*)\}$", source, re.M)
    if not hit:
        raise ValueError(f"Missing macro {name} in {DATA_TEX}")
    return hit.group(1)


def coords(source, name):
    return [(float(x), float(y)) for x, y in re.findall(
        r"\((-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?)\)", macro(source, name)
    )]


def set_text(shape, text, size=12, color=INK, bold=False, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def textbox(slide, x, y, w, h, text, size=12, color=INK, bold=False,
            align=PP_ALIGN.LEFT, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    return set_text(shape, text, size, color, bold, align)


def card(slide, x, y, w, h, text, fill, size=11, name=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = INK
    shape.line.width = Pt(0.9)
    shape.text_frame.margin_left = Inches(0.06)
    shape.text_frame.margin_right = Inches(0.06)
    shape.text_frame.margin_top = Inches(0.03)
    shape.text_frame.margin_bottom = Inches(0.03)
    return set_text(shape, text, size=size)


def group_box(slide, x, y, w, h, label, line_color, fill_color, name):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.fill.transparency = 82
    shape.line.color.rgb = line_color
    shape.line.transparency = 45
    shape.line.width = Pt(1.0)
    textbox(slide, x + .15, y - .02, w - .3, .28, label, 8.5, INK, True,
            PP_ALIGN.LEFT, name + " label")
    return shape


def arrow(slide, x1, y1, x2, y2, name):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.name = name
    line.line.color.rgb = INK
    line.line.width = Pt(1.5)
    ln = line._element.spPr.get_or_add_ln()
    end = ln.find(qn("a:tailEnd"))
    if end is None:
        end = ln.makeelement(qn("a:tailEnd"), {})
        ln.append(end)
    end.set("type", "triangle")
    end.set("w", "sm")
    end.set("len", "sm")
    return line


def style_chart(chart):
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.name = "Aptos"
    chart.legend.font.size = Pt(8.5)
    chart.chart_style = 10
    xaxis = chart.category_axis
    yaxis = chart.value_axis
    xaxis.minimum_scale = -30
    xaxis.maximum_scale = 30
    xaxis.major_unit = 10
    yaxis.minimum_scale = 0
    yaxis.maximum_scale = 100
    yaxis.major_unit = 25
    for axis in (xaxis, yaxis):
        axis.tick_labels.font.name = "Aptos"
        axis.tick_labels.font.size = Pt(8)
        axis.format.line.color.rgb = INK
        axis.format.line.transparency = 40
        axis.has_major_gridlines = True
        axis.major_gridlines.format.line.color.rgb = RGBColor(216, 222, 228)
    xaxis.has_title = True
    xaxis.axis_title.text_frame.text = "Time from annotated goal (s)"
    yaxis.has_title = True
    yaxis.axis_title.text_frame.text = "Goal score (%)"
    for axis in (xaxis, yaxis):
        axis.axis_title.text_frame.paragraphs[0].font.name = "Aptos"
        axis.axis_title.text_frame.paragraphs[0].font.size = Pt(9)


def hide_legend_entries(chart, indices):
    """Hide utility series while keeping them editable in the chart."""
    legend = chart._chartSpace.chart.legend
    for index in indices:
        entry = legend.makeelement(qn("c:legendEntry"), {})
        idx = entry.makeelement(qn("c:idx"), {})
        idx.set("val", str(index))
        delete = entry.makeelement(qn("c:delete"), {})
        delete.set("val", "1")
        entry.append(idx)
        entry.append(delete)
        legend.append(entry)


def main(out):
    source = open(DATA_TEX, encoding="utf-8").read()
    pv, pa, pf = (coords(source, name) for name in ("fwCoordsV", "fwCoordsA", "fwCoordsF"))
    values = {name: macro(source, name) for name in (
        "fwTime", "fwPv", "fwPa", "fwPf", "fwOffset"
    )}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE

    textbox(slide, .48, .18, 12.35, .4,
            "AGREE: reliability from heterogeneous event evidence",
            19, INK, True, PP_ALIGN.LEFT, "Figure title")
    textbox(slide, .48, .62, 4.1, .3, "(a) Genuine SoccerNet event",
            10.5, INK, True, PP_ALIGN.LEFT, "Panel a title")
    textbox(slide, 4.85, .62, 8.0, .3, "(b) Out-of-league branch evidence",
            10.5, INK, True, PP_ALIGN.LEFT, "Panel b title")

    picture = slide.shapes.add_picture(FRAME, Inches(.48), Inches(.96),
                                       width=Inches(4.05), height=Inches(2.28))
    picture.name = "Genuine SoccerNet broadcast frame"
    label = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(.63), Inches(2.83), Inches(1.72), Inches(.28))
    label.name = "Event label background"
    label.fill.solid(); label.fill.fore_color.rgb = INK; label.fill.transparency = 12
    label.line.fill.background()
    set_text(label, "Goal at " + values["fwTime"], 8.5, WHITE, True, PP_ALIGN.CENTER)

    data = XyChartData()
    for name, points in (("Visual Pv", pv), ("Audio Pa", pa), ("Fused pf", pf),
                         ("Operating threshold", [(-30, 40), (30, 40)]),
                         ("Annotation", [(0, 0), (0, 100)])):
        series = data.add_series(name)
        for x, y in points:
            series.add_data_point(x, y)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
        Inches(4.78), Inches(.88), Inches(8.08), Inches(2.48), data
    )
    chart_shape.name = "Editable branch score chart"
    chart = chart_shape.chart
    style_chart(chart)
    hide_legend_entries(chart, (3, 4))
    colors = (VISUAL, AUDIO, INK, RGBColor(118, 127, 137), RGBColor(150, 158, 166))
    widths = (2.25, 2.25, 2.0, 1.0, 1.0)
    for series, color, width in zip(chart.series, colors, widths):
        series.format.line.color.rgb = color
        series.format.line.width = Pt(width)
    chart.series[2].format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    chart.series[3].format.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
    chart.series[4].format.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(.48), Inches(3.50), Inches(12.38), Inches(.66))
    callout.name = "Observed disagreement callout"
    callout.fill.solid(); callout.fill.fore_color.rgb = RGBColor(253, 244, 239)
    callout.line.color.rgb = AUDIO; callout.line.transparency = 45
    set_text(
        callout,
        "SCOPE-DEFINING MISS   Around the fused local peak at %s s, audio reaches pₐ=%s but vision only pᵥ=%s; "
        "the fused peak is p_f=%s < τ₀=0.4. No candidate is emitted, so AGREE cannot assess it."
        % (values["fwOffset"], values["fwPa"], values["fwPv"], values["fwPf"]),
        10.5, INK, False, PP_ALIGN.LEFT,
    )

    textbox(slide, .48, 4.28, 5.2, .27, "(c) Candidate-level reliability workflow",
            10.5, INK, True, PP_ALIGN.LEFT, "Panel c title")
    group_box(slide, .48, 4.62, 4.0, 2.15, "BASE EVENT SPOTTER",
              INK, RGBColor(231, 235, 239), "Base spotter group")
    group_box(slide, 4.66, 4.62, 8.20, 2.15, "AGREE: POST-HOC RELIABILITY LAYER",
              GOLD, RGBColor(255, 247, 224), "AGREE group")

    card(slide, .72, 5.03, 1.35, .60, "Visual\nPv(t,c)", RGBColor(235, 246, 252), 10, "Visual branch")
    card(slide, .72, 5.85, 1.35, .60, "Audio\nPa(t,c)", RGBColor(253, 239, 233), 10, "Audio branch")
    card(slide, 2.54, 5.28, 1.58, .96, "Late fusion + NMS\nCandidate if pf >= tau0",
         RGBColor(255, 255, 255), 9.5, "Candidate gate")
    card(slide, 4.94, 5.19, 2.13, 1.12,
         "Retain separate evidence\nBranch scores + agreement\n+/- 5 s context",
         RGBColor(255, 247, 224), 9.5, "Evidence features")
    card(slide, 7.54, 5.19, 2.18, 1.12,
         "Estimate reliability\nClass-specific g(z) + isotonic\ncalibration -> r(e)",
         RGBColor(255, 247, 224), 9.5, "Reliability estimator")
    card(slide, 10.35, 5.03, 2.02, .60, "Abstain: review if r < theta",
         RGBColor(233, 249, 241), 9.5, "Abstention action")
    card(slide, 10.35, 5.85, 2.02, .60, "Rerank: preserve fused scores",
         RGBColor(233, 249, 241), 9.5, "Reranking action")

    arrow(slide, 2.07, 5.33, 2.54, 5.60, "Visual to candidate")
    arrow(slide, 2.07, 6.15, 2.54, 5.92, "Audio to candidate")
    arrow(slide, 4.12, 5.76, 4.94, 5.76, "Candidates to features")
    arrow(slide, 7.07, 5.76, 7.54, 5.76, "Features to reliability")
    arrow(slide, 9.72, 5.76, 10.35, 5.33, "Reliability to abstention")
    arrow(slide, 9.72, 5.76, 10.35, 6.15, "Reliability to reranking")

    textbox(slide, .48, 7.02, 12.3, .25,
            "Source: genuine SoccerNet broadcast frame; score traces from leave-Premier-League-out seed-0 models. "
            "All elements except the source frame are editable PowerPoint objects.",
            8, RGBColor(92, 105, 119), False, PP_ALIGN.LEFT, "Provenance note")

    prs.save(out)
    print("wrote", out, "| points:", len(pv), len(pa), len(pf))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(MANUSCRIPT, "fig1_editable.pptx"))
